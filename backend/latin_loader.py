"""Latin Quarter data loader.

Ingests Latin Quarter's raw Myntra data (the attached RAR, already extracted to
LQ_DATA_DIR) into the platform's collections, using the *existing* system shapes:

  - Myntra order CSVs  -> `sales`  + `calculations` (Latin adapter: uses Myntra's own
                          reported commission / fixed fee / logistics / tax / settlement
                          as the EXPECTED figures, since LQ's contract masters are partial)
  - Myntra payout CSVs -> `settlement` (actual Settled_Amount, aggregated per order line)
  - Commission masters xlsx -> commission_rules / fixed_fees / subcat_levels (for /masters)

Join key across sales & settlement = (order_release_id, order_line_id).

Run:  python -m latin_loader           (from /app/backend)
"""
import os
import csv
import glob
import uuid
import asyncio
from datetime import datetime, timezone
from pathlib import Path

from dotenv import load_dotenv
load_dotenv(Path(__file__).parent / ".env")

import openpyxl
from motor.motor_asyncio import AsyncIOMotorClient

LQ_DATA_DIR = os.environ.get(
    "LQ_DATA_DIR",
    "/app/lq_data/myntra_unar/Myntra Raw data",
)
MASTERS_XLSX = os.environ.get("LQ_MASTERS_XLSX", "/app/lq_data/commission-masters.xlsx")

client = AsyncIOMotorClient(os.environ["MONGO_URL"])
db = client[os.environ["DB_NAME"]]


def _uid():
    return str(uuid.uuid4())


def _iso():
    return datetime.now(timezone.utc).isoformat()


def _num(v):
    if v is None or v == "":
        return 0.0
    try:
        return float(str(v).replace(",", "").strip())
    except Exception:
        return 0.0


def _s(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _month(*vals):
    for v in vals:
        s = _s(v)
        if not s:
            continue
        # ISO-ish 'YYYY-MM-DD ...'
        if len(s) >= 7 and s[4] == "-":
            return s[:7]
    return None


ZONE_MAP = {"local": "Local", "zonal": "Zonal", "national": "National",
            "l": "Local", "z": "Zonal", "n": "National"}


def _zone(v):
    s = (_s(v) or "").lower()
    return ZONE_MAP.get(s)


# ---------------- Commission masters ----------------
async def load_masters():
    wb = openpyxl.load_workbook(MASTERS_XLSX, read_only=True, data_only=True)

    # Commission Rules: commission % lives in the 'price_range' column (whole number).
    rules = []
    ws = wb["Commission Rules"]
    for row in list(ws.iter_rows(values_only=True))[1:]:
        # (id, brand, master_category, sub_category, gender, lower, upper, pct, model, pct2, active)
        if not row or row[3] is None:
            continue
        lower = _num(row[5]); upper = _num(row[6])
        pct_raw = row[7] if row[7] is not None else row[9]
        pct = _num(pct_raw)
        pct = pct / 100.0 if pct > 1 else pct
        rules.append({
            "id": _uid(), "brand": _s(row[1]) or "Latin Quarter",
            "master_category": (_s(row[2]) or "APPAREL").upper(),
            "sub_category": _s(row[3]), "gender": _s(row[4]) or "Women",
            "lower_limit": lower, "upper_limit": upper,
            "price_range": f"{int(lower)}-{int(upper)}",
            "commission_model": "Latin Quarter Contract",
            "commission_pct": round(pct, 4), "active": True,
        })

    # Fixed Fee: per sub-category slabs (label = sub-category).
    fixed = []
    ws = wb["Fixed Fee"]
    for row in list(ws.iter_rows(values_only=True))[1:]:
        if not row or row[3] is None:
            continue
        fixed.append({
            "id": _uid(), "aisp_lower": _num(row[1]), "aisp_upper": _num(row[2]),
            "label": _s(row[3]), "sub_category": _s(row[3]),
            "fixed_fee": _num(row[4]), "active": True,
        })

    # GT Charges sheet actually carries sub-category -> level mapping.
    levels = []
    ws = wb["GT Charges"]
    for row in list(ws.iter_rows(values_only=True))[1:]:
        sc = _s(row[0]); lvl = _s(row[1])
        if not sc or not lvl or sc.lower() in ("article_type", "sub_category"):
            continue
        levels.append({"id": _uid(), "sub_category": sc, "level": lvl})

    await db.commission_rules.delete_many({})
    await db.fixed_fees.delete_many({})
    await db.subcat_levels.delete_many({})
    if rules:
        await db.commission_rules.insert_many(rules)
    if fixed:
        await db.fixed_fees.insert_many(fixed)
    if levels:
        await db.subcat_levels.insert_many(levels)
    print(f"[masters] commission_rules={len(rules)} fixed_fees={len(fixed)} subcat_levels={len(levels)}")


# ---------------- Orders -> sales + calculations ----------------
def _order_to_docs(row, upload_id):
    orl = _s(row.get("order_release_id"))
    oli = _s(row.get("order_line_id"))
    if not orl or not oli:
        return None
    return_type = _s(row.get("return_type"))
    is_return = bool(return_type) and return_type.lower() not in ("null", "none", "0", "false")
    if not is_return:
        return_type = None
    nsv = _num(row.get("seller_product_amount"))
    nsv_val = round(abs(nsv), 2)
    sub_category = _s(row.get("article_type"))
    zone = _zone(row.get("shipment_zone_classification"))
    rm = _month(row.get("delivery_date"), row.get("packing_date"),
                row.get("return_date"), row.get("order_created_date"))
    sales_id = _uid()

    commission = _num(row.get("total_commission"))
    fixed_fee = _num(row.get("fixed_fee"))
    gt = _num(row.get("total_logistics_deduction"))
    tcs = _num(row.get("tcs_amount"))
    tds = _num(row.get("tds_amount"))
    settlement = round(_num(row.get("total_settlement")), 2)
    pct = _num(row.get("commission_percentage"))
    pct = pct / 100.0 if pct > 1 else pct
    nsv_after_gt = round(nsv_val - gt, 2)
    total_deductions = round(commission + fixed_fee + gt + tcs + tds, 2)
    order_type = "return" if is_return else "sales"

    sales_doc = {
        "id": sales_id, "upload_id": upload_id, "uploaded_at": _iso(),
        "source_file": "Myntra Data (Latin Quarter)",
        "online_order_id": orl, "order_line_id": oli, "sku": _s(row.get("sku_code")),
        "sub_category": sub_category, "category": sub_category, "main_category": "APPAREL",
        "brand": _s(row.get("brand")), "gender": _s(row.get("gender")),
        "zone": zone or _s(row.get("shipment_zone_classification")),
        "qty": 1.0, "mrp": _num(row.get("mrp")),
        "nsv_val": nsv_val, "nsv_per_unit": nsv_val,
        "txn_type": "Return" if is_return else "Sales",
        "order_status": return_type or "Delivered",
        "portal_name": "Myntra", "report_month": rm,
        "order_date": _s(row.get("packing_date")), "posting_date": _s(row.get("delivery_date")),
        "month": rm,
        "actual_commission_value": commission, "actual_fixed_fee": fixed_fee,
        "actual_gt_amount": gt, "actual_return_fee": 0.0,
        "myntra_total_settlement": settlement,
        "myntra_actual_settlement": round(_num(row.get("total_actual_settlement")), 2),
        "amount_pending_settlement": round(_num(row.get("amount_pending_settlement")), 2),
    }

    calc_doc = {
        "id": _uid(), "sales_id": sales_id, "upload_id": upload_id,
        "online_order_id": orl, "order_line_id": oli, "sku": _s(row.get("sku_code")),
        "computed_at": _iso(),
        "commission_base": round(commission / 1.18, 2), "commission_gst": round(commission - commission / 1.18, 2),
        "commission_incl_gst": round(commission, 2), "commission_pct": round(pct, 4) if pct else None,
        "fixed_fee": round(fixed_fee, 2), "fixed_fee_gst": 0.0, "fixed_fee_incl_gst": round(fixed_fee, 2),
        "gt_charge": round(gt, 2), "return_fee": 0.0,
        "tcs": round(tcs, 2), "tds": round(tds, 2),
        "nsv_after_gt": nsv_after_gt,
        "total_deductions": total_deductions, "expected_settlement": settlement,
        "order_type": order_type, "is_return": is_return,
        "unmapped": False, "unmapped_reasons": [],
        "report_month": rm,
        "breakdown": {
            "isp": abs(nsv_val), "qty": 1.0, "nsv_val": nsv_val,
            "sub_category": sub_category, "master_category": "APPAREL",
            "zone": zone, "level": _s(row.get("article_level")),
            "order_type": order_type, "report_month": rm, "nsv_after_gt": nsv_after_gt,
            "commission_rule": {"commission_pct": round(pct, 4) if pct else None},
            "fixed_fee_slab": {"label": sub_category, "fixed_fee": round(fixed_fee, 2)},
            "gt_charge_cell": {"unit_charge": round(gt, 2), "qty": 1.0},
            "return_fee_cell": {"fee": 0.0, "applied": False},
        },
    }
    return sales_doc, calc_doc


async def load_orders():
    upload_id = _uid()
    files = sorted(glob.glob(os.path.join(LQ_DATA_DIR, "Myntra Data", "*.csv")))
    sales_batch, calc_batch = [], []
    n_sales = 0
    months = {}
    await db.sales.delete_many({})
    await db.calculations.delete_many({})
    for fp in files:
        with open(fp, encoding="utf-8", errors="replace", newline="") as fh:
            reader = csv.DictReader(fh)
            for row in reader:
                out = _order_to_docs(row, upload_id)
                if not out:
                    continue
                sd, cd = out
                sales_batch.append(sd)
                calc_batch.append(cd)
                if sd.get("report_month"):
                    months[sd["report_month"]] = months.get(sd["report_month"], 0) + 1
                if len(sales_batch) >= 1000:
                    await db.sales.insert_many(sales_batch)
                    await db.calculations.insert_many(calc_batch)
                    n_sales += len(sales_batch)
                    sales_batch, calc_batch = [], []
    if sales_batch:
        await db.sales.insert_many(sales_batch)
        await db.calculations.insert_many(calc_batch)
        n_sales += len(sales_batch)

    await db.uploads.delete_many({"type": "sales"})
    await db.uploads.insert_one({
        "id": upload_id, "type": "sales", "filename": "Myntra Order Data (Latin Quarter)",
        "uploaded_at": _iso(), "sheet": "Myntra Data",
        "accepted_count": n_sales, "rejected_count": 0, "rejections_sample": [],
        "months": months, "status": "processed",
    })
    print(f"[orders] sales+calc rows={n_sales} months={months} files={len(files)}")
    return n_sales


# ---------------- Payouts -> settlement ----------------
async def load_payouts():
    upload_id = _uid()
    files = sorted(glob.glob(os.path.join(LQ_DATA_DIR, "Payouts", "*", "*.csv")))
    agg = {}
    for fp in files:
        with open(fp, encoding="utf-8", errors="replace", newline="") as fh:
            reader = csv.DictReader(fh)
            for row in reader:
                orl = _s(row.get("order_release_id"))
                oli = _s(row.get("order_line_id"))
                if not orl or not oli:
                    continue
                key = (orl, oli)
                a = agg.setdefault(key, {
                    "online_order_id": orl, "order_line_id": oli, "sku": _s(row.get("Store_Order_id")),
                    "settled_commission": 0.0, "settled_fixed_fee": 0.0, "settled_gt_charge": 0.0,
                    "settled_return_fee": 0.0, "settled_tcs": 0.0, "settled_tds": 0.0,
                    "settled_amount": 0.0, "settlement_date": None, "report_month": None,
                    "payment_type": _s(row.get("Payment_Type")),
                })
                a["settled_commission"] += _num(row.get("Commission"))
                a["settled_fixed_fee"] += _num(row.get("fixed_fee"))
                a["settled_gt_charge"] += _num(row.get("Logistics_Commission"))
                a["settled_tds"] += _num(row.get("TDS"))
                a["settled_amount"] += _num(row.get("Settled_Amount"))
                pd = _s(row.get("Payment_Date"))
                if pd:
                    a["settlement_date"] = pd
                    a["report_month"] = a["report_month"] or (pd[:7] if len(pd) >= 7 else None)

    docs = []
    for a in agg.values():
        for k in ("settled_commission", "settled_fixed_fee", "settled_gt_charge",
                  "settled_return_fee", "settled_tcs", "settled_tds", "settled_amount"):
            a[k] = round(a[k], 2)
        a.update({"id": _uid(), "upload_id": upload_id, "uploaded_at": _iso(),
                  "source_file": "Myntra Payouts (Latin Quarter)"})
        docs.append(a)

    await db.settlement.delete_many({})
    for i in range(0, len(docs), 1000):
        await db.settlement.insert_many(docs[i:i + 1000])

    await db.uploads.delete_many({"type": "settlement"})
    await db.uploads.insert_one({
        "id": upload_id, "type": "settlement", "filename": "Myntra Payouts (Latin Quarter)",
        "uploaded_at": _iso(), "sheet": "Payouts",
        "accepted_count": len(docs), "rejected_count": 0, "rejections_sample": [], "status": "processed",
    })
    print(f"[payouts] settlement lines={len(docs)} files={len(files)}")
    return len(docs)


async def main():
    print("=== Latin Quarter loader start ===")
    await load_masters()
    n_sales = await load_orders()
    n_settle = await load_payouts()
    # clear stale recon artefacts; recon is triggered via the API after load
    await db.discrepancies.delete_many({})
    await db.recon_runs.delete_many({})
    print(f"=== done: sales={n_sales} settlement={n_settle} ===")


if __name__ == "__main__":
    asyncio.run(main())
